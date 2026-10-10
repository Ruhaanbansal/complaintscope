"""ComplaintScope: Academic Report Builder.
Compiles report/report.md into report/ComplaintScope_Academic_Report.docx (and copies to ComplaintScope_Report.docx).
"""
import re
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def build_report():
    md_path = Path("report/report.md")
    if not md_path.exists():
        print("report/report.md not found.")
        return

    doc = Document()

    # Page Margins: 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    lines = md_path.read_text(encoding="utf-8").splitlines()

    i = 0
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return

        parsed_rows = []
        for tr in table_rows:
            cells = [c.strip() for c in tr.strip('|').split('|')]
            if cells and all(re.match(r"^:?-+:?$", c) for c in cells if c):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            in_table = False
            table_rows = []
            return

        col_count = max(len(r) for r in parsed_rows)
        for r in parsed_rows:
            while len(r) < col_count:
                r.append("")

        t = doc.add_table(rows=len(parsed_rows), cols=col_count)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True

        for r_idx, row in enumerate(parsed_rows):
            is_header = (r_idx == 0)
            for c_idx, val in enumerate(row):
                cell = t.cell(r_idx, c_idx)
                cell.text = val.replace("**", "").replace("*", "")
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(3)
                cp.paragraph_format.space_after = Pt(3)
                run = cp.runs[0] if cp.runs else cp.add_run()
                run.font.name = 'Calibri'
                run.font.size = Pt(10)
                if is_header:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    set_cell_background(cell, "1E3A8A")
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        in_table = False
        table_rows = []

    fig_map = {
        "Fig 1": "results/figures/fig_01_funnel.png",
        "Fig 2": "results/figures/fig_02_monthly_volume.png",
        "Fig 3": "results/figures/fig_03_product_distribution.png",
        "Fig 4": "results/figures/fig_04_narrative_length.png",
        "Fig 5": "results/figures/fig_05_relief_rate_by_product.png",
        "Fig 6": "results/figures/fig_06_split_timeline.png",
        "Fig 7": "results/figures/fig07_zipf_distribution.png",
        "Fig 8": "results/figures/fig08_class_top_terms.png",
        "Fig 9": "results/figures/fig09_2d_projection.png",
        "Fig 10": "results/figures/fig10_confusion_matrix.png",
        "Fig 11": "results/figures/fig11_calibration_curve.png",
        "Fig 12": "results/figures/fig12_pr_curve.png",
        "Fig 13": "results/figures/fig13_triage_gain.png",
        "Fig 14": "results/figures/fig14_lda_coherence_sweep.png",
        "Fig 15": "results/figures/fig15_topic_product_heatmap.png",
        "Fig 16": "results/figures/fig16_topic_burst_detection.png",
        "Fig 17": "results/figures/fig17_company_network.png",
        "Fig 18": "results/figures/fig18_systemic_risk_index.png",
        "Fig 19": "results/figures/fig19_sri_sensitivity.png",
    }

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_rows.append(line.strip())
            i += 1
            continue
        elif in_table:
            flush_table()

        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(24)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith("## "):
            h = doc.add_heading(line[3:].strip(), level=1)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
        elif line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=2)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith("#### "):
            h = doc.add_heading(line[5:].strip(), level=3)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(11)
                r.font.bold = True
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            clean_text = line[2:].strip().replace("**", "").replace("*", "")
            p.text = clean_text
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style='List Number')
            clean_text = re.sub(r"^\d+\.\s", "", line).strip().replace("**", "").replace("*", "")
            p.text = clean_text
            p.paragraph_format.space_after = Pt(2)
        elif line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
        elif line.strip():
            clean_text = line.replace("**", "").replace("*", "")
            p = doc.add_paragraph(clean_text)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15

            for fig_key, fig_file in fig_map.items():
                if fig_key in line and Path(fig_file).exists():
                    try:
                        doc.add_picture(fig_file, width=Inches(5.5))
                        last_p = doc.paragraphs[-1]
                        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap = doc.add_paragraph(f"{fig_key}: Embedded from {fig_file}")
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap.runs[0].font.size = Pt(9)
                        cap.runs[0].font.italic = True
                        cap.runs[0].font.color.rgb = RGBColor(100, 116, 139)
                    except Exception as e:
                        print(f"Failed to embed {fig_file}: {e}")

        i += 1

    if in_table:
        flush_table()

    out_file = Path("report/ComplaintScope_Academic_Report.docx")
    doc.save(str(out_file))
    print(f"Report DOCX saved to {out_file}")

    # Copy to ComplaintScope_Report.docx
    dst_file = Path("report/ComplaintScope_Report.docx")
    shutil.copy2(out_file, dst_file)
    print(f"Also copied to {dst_file}")


def main():
    print("Building academic report DOCX...")
    build_report()


if __name__ == "__main__":
    main()
